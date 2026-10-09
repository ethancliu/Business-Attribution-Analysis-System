# -*- coding: utf-8 -*-
"""附件服务：上传落盘、类型识别、轻量解析、列表与删除。

存储规范（需求 §十）: uploads/{user_id}/{conversation_id}/{时间戳}_{原文件名}
"""

from __future__ import annotations

import csv
import io
import json
import logging
from datetime import datetime
from pathlib import Path

from fastapi import HTTPException, status as http_status
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.paths import (
    KIND_UPLOAD,
    abs_path,
    delete_file,
    dir_of,
    human_size,
    safe_join,
    sanitize_filename,
    unique_name,
)
from app.models import Attachment, User

logger = logging.getLogger(__name__)

# 允许上传的扩展名 -> 归一化类型
EXT_TYPE = {
    ".csv": "csv",
    ".xlsx": "excel",
    ".xls": "excel",
    ".json": "json",
    ".txt": "text",
    ".md": "text",
    ".log": "text",
    ".sql": "text",
    ".pdf": "pdf",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".gif": "image",
    ".webp": "image",
    ".zip": "archive",
}

TEXT_TYPES = {"csv", "json", "text"}

PARSE_STATUS_TEXT = {
    "pending": "待解析",
    "parsing": "解析中",
    "parsed": "已解析",
    "failed": "解析失败",
}

PREVIEW_CHARS = 2000


def detect_type(filename: str) -> str:
    return EXT_TYPE.get(Path(filename).suffix.lower(), "unknown")


def _iso(dt: datetime | None) -> str | None:
    return dt.isoformat() if dt else None


# ---------------------------------------------------------------------------
# 上传
# ---------------------------------------------------------------------------

def save_upload(
    db: Session,
    user: User,
    conversation_id: int,
    filename: str,
    content: bytes,
) -> Attachment:
    """把上传内容写入 uploads 目录并登记附件记录。"""
    file_type = detect_type(filename)
    if file_type == "unknown":
        raise HTTPException(
            status_code=http_status.HTTP_400_BAD_REQUEST,
            detail=f"不支持的文件类型：{Path(filename).suffix or '未知'}",
        )

    safe_name = sanitize_filename(filename, fallback="upload.bin")
    stored_name = unique_name(safe_name)
    target_dir = dir_of(KIND_UPLOAD, user.id, conversation_id)
    target = safe_join(target_dir, stored_name)
    target.write_bytes(content)

    rel = f"{KIND_UPLOAD}/{user.id}/{conversation_id}/{stored_name}"

    att = Attachment(
        conversation_id=conversation_id,
        file_name=safe_name,
        file_path=rel,
        file_type=file_type,
        file_size=len(content),
        parse_status="pending",
    )
    db.add(att)
    db.commit()
    db.refresh(att)
    logger.info(
        "会话 #%s 上传附件 #%s %s（%s）", conversation_id, att.id, safe_name, human_size(len(content))
    )

    # 立即做一次轻量解析，便于前端上传后即可预览
    try:
        parse_attachment(db, att)
    except Exception as exc:  # noqa: BLE001 - 解析失败不应影响上传本身
        logger.warning("附件 #%s 解析失败: %s", att.id, exc)
        att.parse_status = "failed"
        db.commit()

    return att


# ---------------------------------------------------------------------------
# 解析（轻量：只读取结构概要，不引入重型依赖）
# ---------------------------------------------------------------------------

def parse_attachment(db: Session, att: Attachment) -> dict:
    """解析附件结构概要，结果写入 parse 状态。返回概要字典。"""
    att.parse_status = "parsing"
    db.commit()

    path = abs_path(att.file_path)
    if not path.is_file():
        att.parse_status = "failed"
        db.commit()
        return {"error": "文件不存在"}

    summary: dict = {"file_type": att.file_type, "file_size": att.file_size}

    if att.file_type == "csv":
        summary.update(_parse_csv(path))
    elif att.file_type == "excel":
        summary.update(_parse_excel(path))
    elif att.file_type == "json":
        summary.update(_parse_json(path))
    elif att.file_type == "text":
        summary.update(_parse_text(path))
    else:
        summary["note"] = "该类型不做结构化解析，分析时按整文件读取"

    att.parse_status = "parsed"
    db.commit()
    return summary


def _parse_csv(path: Path) -> dict:
    rows = 0
    header: list[str] = []
    sample: list[list[str]] = []
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        reader = csv.reader(fh)
        for i, row in enumerate(reader):
            if i == 0:
                header = [c.strip() for c in row]
            else:
                rows += 1
                if len(sample) < 5:
                    sample.append(row[:20])
            if i > 200000:  # 安全上限，避免超大文件拖垮解析
                break
    return {"row_count": rows, "columns": header, "sample": sample}


def _parse_excel(path: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover
        return {"note": "未安装 openpyxl，跳过解析"}

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = wb[wb.sheetnames[0]]
        header: list[str] = []
        sample: list[list] = []
        rows = 0
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i == 0:
                header = [str(c) if c is not None else "" for c in row]
            else:
                rows += 1
                if len(sample) < 5:
                    sample.append([str(c) if c is not None else "" for c in row[:20]])
            if i > 200000:
                break
        return {
            "row_count": rows,
            "columns": header,
            "sample": sample,
            "sheets": list(wb.sheetnames),
        }
    finally:
        wb.close()


def _parse_json(path: Path) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        return {"note": f"JSON 解析失败：{exc}", "preview": text[:PREVIEW_CHARS]}
    if isinstance(data, list):
        return {"row_count": len(data), "root": "array", "sample": data[:3]}
    if isinstance(data, dict):
        return {"row_count": len(data), "root": "object", "keys": list(data.keys())[:30]}
    return {"row_count": 1, "root": type(data).__name__}


def _parse_text(path: Path) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    lines = text.splitlines()
    return {"line_count": len(lines), "preview": text[:PREVIEW_CHARS]}


def preview_text(att: Attachment, limit: int = PREVIEW_CHARS) -> str:
    """给前端的文本预览。"""
    try:
        path = abs_path(att.file_path)
    except ValueError:
        return ""
    if not path.is_file():
        return ""
    if att.file_type in TEXT_TYPES or att.file_type == "text":
        return path.read_text(encoding="utf-8", errors="replace")[:limit]
    if att.file_type == "excel":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            sheet = wb[wb.sheetnames[0]]
            buf = io.StringIO()
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i > 20:
                    break
                buf.write("\t".join("" if c is None else str(c) for c in row) + "\n")
            wb.close()
            return buf.getvalue()[:limit]
        except Exception as exc:  # noqa: BLE001
            return f"[Excel 预览失败] {exc}"
    return f"[{att.file_type} 类型暂不支持文本预览]"


# ---------------------------------------------------------------------------
# 查询 / 删除
# ---------------------------------------------------------------------------

def get_owned(db: Session, user: User, attachment_id: int) -> Attachment:
    att = db.get(Attachment, attachment_id)
    if att is None:
        raise HTTPException(status_code=http_status.HTTP_404_NOT_FOUND, detail="附件不存在")
    conv = att.conversation
    if conv is None or conv.user_id != user.id:
        raise HTTPException(status_code=http_status.HTTP_404_NOT_FOUND, detail="附件不存在或无权访问")
    return att


def list_for_conversation(db: Session, conversation_id: int) -> list[Attachment]:
    stmt = (
        select(Attachment)
        .where(Attachment.conversation_id == conversation_id)
        .order_by(Attachment.id.desc())
    )
    return list(db.execute(stmt).scalars().all())


def brief(att: Attachment) -> dict:
    return {
        "attachment_id": att.id,
        "file_name": att.file_name,
        "file_type": att.file_type,
        "file_size": att.file_size,
        "file_size_text": human_size(att.file_size),
        "parse_status": att.parse_status,
        "created_at": _iso(att.created_at),
    }


def item(att: Attachment) -> dict:
    return {
        **brief(att),
        "parse_status_text": PARSE_STATUS_TEXT.get(att.parse_status, att.parse_status),
        "message_id": att.message_id,
    }


def detail(att: Attachment) -> dict:
    return {
        "attachment_id": att.id,
        "conversation_id": att.conversation_id,
        "message_id": att.message_id,
        "file_name": att.file_name,
        "file_path": att.file_path,
        "file_type": att.file_type,
        "file_size": att.file_size,
        "parse_status": att.parse_status,
        "created_at": _iso(att.created_at),
        "preview": preview_text(att, limit=500),
    }


def delete_many(db: Session, user: User, attachment_ids: list[int]) -> dict:
    """删除附件记录与物理文件。"""
    wanted = list(dict.fromkeys(attachment_ids))
    if not wanted:
        return {"deleted": 0, "attachment_ids": [], "removed_files": 0}

    atts = db.execute(select(Attachment).where(Attachment.id.in_(wanted))).scalars().all()
    owned: list[Attachment] = []
    for a in atts:
        conv = a.conversation
        if conv is not None and conv.user_id == user.id:
            owned.append(a)

    removed = 0
    for a in owned:
        if delete_file(a.file_path):
            removed += 1
        db.delete(a)
    db.commit()

    ids = [a.id for a in owned]
    logger.info("用户 %s 删除附件 %s（物理文件 %s 个）", user.username, ids, removed)
    return {"deleted": len(owned), "attachment_ids": ids, "removed_files": removed}


def attachments_for_tasks(db: Session, conversation_id: int) -> list[Attachment]:
    """取会话内所有已解析附件，供分析引擎做「基于文件的分析」。"""
    stmt = select(Attachment).where(
        Attachment.conversation_id == conversation_id,
        Attachment.parse_status == "parsed",
    ).order_by(Attachment.id)
    return list(db.execute(stmt).scalars().all())
