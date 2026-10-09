# -*- coding: utf-8 -*-
"""文件读取工具：读取用户上传附件或分析工作区文件，产出证据。

安全约束:
    只能读取 data/ 根目录之内的文件（路径穿越一律拒绝），
    读取内容有长度上限，避免把超大文件灌进上下文。
"""

from __future__ import annotations

import csv
import io
import json
import logging
from pathlib import Path
from typing import Any

from app.analysis.tools.base import Tool, ToolContext, ToolResult, register
from app.core.paths import abs_path, human_size

logger = logging.getLogger(__name__)

MAX_CHARS = 20000
MAX_PREVIEW_ROWS = 30


def _read_text(path: Path, limit: int = MAX_CHARS) -> str:
    return path.read_text(encoding="utf-8", errors="replace")[:limit]


def _read_csv(path: Path) -> dict:
    rows: list[list[str]] = []
    header: list[str] = []
    total = 0
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        reader = csv.reader(fh)
        for i, row in enumerate(reader):
            if i == 0:
                header = [c.strip() for c in row]
            else:
                total += 1
                if len(rows) < MAX_PREVIEW_ROWS:
                    rows.append(row[:20])
            if i > 500000:
                break
    return {"header": header, "row_count": total, "preview": rows}


def _read_excel(path: Path) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = wb[wb.sheetnames[0]]
        header: list[str] = []
        rows: list[list] = []
        total = 0
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i == 0:
                header = [str(c) if c is not None else "" for c in row]
            else:
                total += 1
                if len(rows) < MAX_PREVIEW_ROWS:
                    rows.append([str(c) if c is not None else "" for c in row[:20]])
            if i > 200000:
                break
        return {"header": header, "row_count": total, "preview": rows, "sheets": list(wb.sheetnames)}
    finally:
        wb.close()


class FileReadTool(Tool):
    name = "file_read"
    description = "读取用户上传的附件或工作区文件内容，支持 csv/excel/json/文本，用作分析证据"
    signature = {
        "file_path": "相对 data/ 的路径，形如 uploads/2/12/xxx.csv",
        "mode": "auto | text | csv | excel | json",
    }

    def run(self, ctx: ToolContext, file_path: str = "", mode: str = "auto", **_: Any) -> ToolResult:
        if not file_path:
            return ToolResult.failure("file_path 不能为空")

        try:
            path = abs_path(file_path)
        except ValueError as exc:
            return ToolResult.failure(f"拒绝读取越权路径：{exc}")

        if not path.is_file():
            return ToolResult.failure(f"文件不存在：{file_path}")

        suffix = path.suffix.lower()
        kind = mode if mode != "auto" else {
            ".csv": "csv", ".xlsx": "excel", ".xls": "excel",
            ".json": "json", ".txt": "text", ".md": "text", ".log": "text", ".sql": "text",
        }.get(suffix, "text")

        ctx.check_cancel()

        try:
            if kind == "csv":
                data = _read_csv(path)
                summary = (
                    f"CSV 共 {data['row_count']} 行，字段 {len(data['header'])} 个："
                    f"{', '.join(data['header'][:8])}"
                )
                rows = data["row_count"]
            elif kind == "excel":
                data = _read_excel(path)
                summary = (
                    f"Excel 工作表 {data.get('sheets', [])}，共 {data['row_count']} 行，"
                    f"字段：{', '.join(data['header'][:8])}"
                )
                rows = data["row_count"]
            elif kind == "json":
                payload = json.loads(_read_text(path))
                if isinstance(payload, list):
                    data = {"root": "array", "row_count": len(payload), "preview": payload[:MAX_PREVIEW_ROWS]}
                else:
                    data = {"root": "object", "keys": list(payload.keys())[:30]}
                summary = f"JSON 结构：{data.get('root')}，规模 {data.get('row_count', len(data.get('keys', [])))}"
                rows = int(data.get("row_count", 0))
            else:
                content = _read_text(path)
                data = {"text": content, "chars": len(content)}
                summary = f"文本文件 {path.name}（{human_size(path.stat().st_size)}），已读取 {len(content)} 字符"
                rows = len(content.splitlines())
        except Exception as exc:  # noqa: BLE001
            return ToolResult.failure(f"读取文件失败：{exc}")

        return ToolResult(ok=True, summary=summary, data=data, rows=rows)


register(FileReadTool())
